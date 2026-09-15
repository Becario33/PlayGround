# PlayGround / whatsapp: consulta SQL → WhatsApp Web (grupo Data Analytics).
#   python main.py
# Primera vez: QR. Cierra Chrome antes de volver a correr.

import argparse
import os
import sys
import time

URL_WA = "https://web.whatsapp.com/"
PERFIL = "chrome_whatsapp_perfil"
GRUPO_DEFAULT = "Data Analytics"
TEXTO_DEFAULT = "Prueba de persistencia: la sesión sigue abierta; el QR solo se escanea una vez."
CONSULTA = """
SELECT
    FechaComscore,
    SUM(CAST(Asistencia AS bigint)) AS AsistenciaIndustria
FROM [Programacion].[dbo].[ComscoreMPAMexico]
WHERE FechaComscore = DATEADD(day, -1, CAST(GETDATE() AS date))
GROUP BY FechaComscore
""".strip()
DRIVERS_ODBC = (
    "ODBC Driver 17 for SQL Server",
    "ODBC Driver 18 for SQL Server",
    "SQL Server",
)


def _dir_script():
    try:
        return os.path.dirname(os.path.abspath(__file__))
    except NameError:
        return os.getcwd()


def _cargar_env():
    ruta = os.path.join(_dir_script(), ".env")
    if not os.path.isfile(ruta):
        return False
    with open(ruta, encoding="utf-8") as f:
        for linea in f:
            linea = linea.strip()
            if not linea or linea.startswith("#") or "=" not in linea:
                continue
            k, v = linea.split("=", 1)
            os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))
    return True


def _conectar_sql():
    try:
        import pyodbc
    except ImportError as e:
        raise RuntimeError(
            "Falta pyodbc. En whatsapp:\n"
            "  python -m pip install -r requirements.txt"
        ) from e
    server = os.environ.get("SQL_SERVER", "").strip()
    database = os.environ.get("SQL_DATABASE", "").strip()
    user = os.environ.get("SQL_USER", "").strip()
    password = os.environ.get("SQL_PASSWORD", "")
    if not server or not database or not user:
        raise RuntimeError(
            "Falta .env (SQL_SERVER, SQL_DATABASE, SQL_USER, SQL_PASSWORD).\n"
            "  copy .env.example .env"
        )
    ultimo = None
    for driver in DRIVERS_ODBC:
        extra = "TrustServerCertificate=yes;Encrypt=no;" if "18" in driver else ""
        conn_str = (
            f"DRIVER={{{driver}}};SERVER={server};DATABASE={database};"
            f"UID={user};PWD={password};{extra}"
        )
        try:
            return pyodbc.connect(conn_str, timeout=30)
        except Exception as e:
            ultimo = e
    raise RuntimeError(
        "No conectó a SQL Server (¿estás en el corporativo?). "
        + str(ultimo).split("\n")[0]
    )


def _celda(valor):
    if valor is None:
        return "NULL"
    if hasattr(valor, "strftime"):
        return valor.strftime("%Y-%m-%d")
    return str(valor)


def consulta_asistencia():
    conn = _conectar_sql()
    try:
        cur = conn.cursor()
        cur.execute(CONSULTA)
        cols = [d[0] for d in cur.description]
        filas = [tuple(f) for f in cur.fetchall()]
    finally:
        conn.close()
    lineas = [" | ".join(cols)]
    if not filas:
        lineas.append("(sin filas para ayer)")
    else:
        for fila in filas:
            lineas.append(" | ".join(_celda(v) for v in fila))
    return "\n".join(lineas), cols, filas


def _dir_resultados():
    ruta = os.path.join(_dir_script(), "resultados")
    os.makedirs(ruta, exist_ok=True)
    return ruta


def armar_excel(cols, filas):
    from datetime import datetime
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Comscore"
    ws["A1"] = "Asistencia industria Comscore (ayer)"
    ws["A1"].font = Font(bold=True, size=14)
    n = max(len(cols), 1)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    encabezado = PatternFill("solid", fgColor="217346")
    fuente = Font(bold=True, color="FFFFFF")
    borde = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )
    for i, col in enumerate(cols, 1):
        cel = ws.cell(2, i, col)
        cel.fill = encabezado
        cel.font = fuente
        cel.alignment = Alignment(horizontal="center")
        cel.border = borde
    if not filas:
        ws.cell(3, 1, "(sin filas para ayer)").border = borde
        ws.cell(3, 2, "").border = borde
    else:
        for r, fila in enumerate(filas, 3):
            for c, valor in enumerate(fila, 1):
                cel = ws.cell(r, c)
                if hasattr(valor, "strftime"):
                    cel.value = valor.strftime("%Y-%m-%d")
                    cel.alignment = Alignment(horizontal="center")
                else:
                    cel.value = valor
                    if isinstance(valor, (int, float)):
                        cel.number_format = "#,##0"
                        cel.alignment = Alignment(horizontal="right")
                cel.border = borde
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 24
    nombre = datetime.now().strftime("asistencia_%Y-%m-%d.xlsx")
    ruta = os.path.join(_dir_resultados(), nombre)
    wb.save(ruta)
    return ruta


def _fuente_tabla(escala):
    from PIL import ImageFont

    candidatos = (
        ("arial.ttf", "arialbd.ttf"),
        ("Arial.ttf", "Arial Bold.ttf"),
        ("/Windows/Fonts/arial.ttf", "/Windows/Fonts/arialbd.ttf"),
        ("/System/Library/Fonts/Supplemental/Arial.ttf",
         "/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
    )
    for regular, bold in candidatos:
        try:
            return (
                ImageFont.truetype(regular, 16 * escala),
                ImageFont.truetype(bold, 18 * escala),
                ImageFont.truetype(bold, 16 * escala),
            )
        except Exception:
            continue
    font = ImageFont.load_default()
    return font, font, font


def _captura_pillow(xlsx_path, jpg_path):
    from openpyxl import load_workbook
    from PIL import Image, ImageDraw

    wb = load_workbook(xlsx_path)
    ws = wb.active
    n_cols = max(ws.max_column or 1, 1)
    filas = []
    for row in ws.iter_rows(min_row=1, max_col=n_cols, max_row=ws.max_row, values_only=True):
        celdas = []
        for v in row:
            if v is None:
                celdas.append("")
            elif isinstance(v, (int, float)) and not isinstance(v, bool):
                celdas.append(f"{int(v):,}")
            else:
                celdas.append(str(v))
        if any(c.strip() for c in celdas):
            filas.append(celdas)
    if not filas:
        filas = [["(sin filas)"]]

    escala = 3
    pad_x, pad_y = 10 * escala, 3 * escala
    font, font_t, font_b = _fuente_tabla(escala)
    medidor = ImageDraw.Draw(Image.new("RGB", (8, 8)))

    def _font_fila(i):
        if i == 0:
            return font_t
        if i == 1:
            return font_b
        return font

    def _recorte_texto(txt, fnt, fill, bg):
        from PIL import ImageChops

        txt = txt or " "
        b = medidor.textbbox((0, 0), txt, font=fnt)
        w = max(int(b[2] - b[0]) + 80, 80)
        h = max(int(b[3] - b[1]) + 80, 80)
        im = Image.new("RGB", (w, h), bg)
        ImageDraw.Draw(im).text((40, 40), txt, font=fnt, fill=fill)
        diff = ImageChops.difference(im, Image.new("RGB", im.size, bg))
        mask = diff.convert("L").point(lambda p: 255 if p > 18 else 0)
        box = mask.getbbox()
        if not box:
            return Image.new("RGB", (1, 1), bg)
        return im.crop(box)

    ncols = max((len(f) for f in filas), default=1)
    recortes = []
    for i, fila in enumerate(filas):
        fnt = _font_fila(i)
        if i == 0:
            recortes.append([_recorte_texto(fila[0], fnt, "#217346", "white")])
        elif i == 1:
            recortes.append([
                _recorte_texto(fila[j] if j < len(fila) else "", fnt, "white", "#217346")
                for j in range(ncols)
            ])
        else:
            recortes.append([
                _recorte_texto(fila[j] if j < len(fila) else "", fnt, "#222222", "white")
                for j in range(ncols)
            ])

    anchos = []
    for c in range(ncols):
        mx = 1
        for i, fila_r in enumerate(recortes):
            if i == 0 or c >= len(fila_r):
                continue
            mx = max(mx, fila_r[c].size[0] + pad_x * 2)
        anchos.append(mx)
    titulo_w = recortes[0][0].size[0] + pad_x * 2
    if titulo_w > sum(anchos):
        anchos[-1] += titulo_w - sum(anchos)

    altos = [max(im.size[1] for im in fila_r) + pad_y * 2 for fila_r in recortes]
    w, h = sum(anchos), sum(altos)
    tabla = Image.new("RGB", (w, h), "white")
    y = 0
    for i, fila_r in enumerate(recortes):
        ah = altos[i]
        x = 0
        if i == 0:
            tabla.paste(fila_r[0], (pad_x, y + (ah - fila_r[0].size[1]) // 2))
        elif i == 1:
            tabla.paste(Image.new("RGB", (w, ah), "#217346"), (0, y))
            for j, im in enumerate(fila_r):
                tabla.paste(im, (x + pad_x, y + (ah - im.size[1]) // 2))
                x += anchos[j]
        else:
            for j, im in enumerate(fila_r):
                tabla.paste(im, (x + pad_x, y + (ah - im.size[1]) // 2))
                x += anchos[j]
        y += ah

    draw = ImageDraw.Draw(tabla)
    draw.rectangle([0, 0, w - 1, h - 1], outline="#C8C8C8")
    y = 0
    for ah in altos[:-1]:
        y += ah
        draw.line([(0, y), (w - 1, y)], fill="#C8C8C8")
    x = 0
    for c in range(ncols - 1):
        x += anchos[c]
        draw.line([(x, altos[0]), (x, h - 1)], fill="#C8C8C8")
    # Solo la tabla (sin lienzo gris). El anti-sticker ya va por paste/Document.
    tabla = tabla.convert("RGB")
    tabla.save(jpg_path, format="JPEG", quality=95, optimize=True)
    print("  JPEG foto:", tabla.size[0], "x", tabla.size[1])
    return jpg_path


def captura_excel(xlsx_path):
    jpg_path = os.path.splitext(xlsx_path)[0] + ".jpg"
    _captura_pillow(xlsx_path, jpg_path)
    print("Imagen (JPEG, no sticker):")
    print(" ", jpg_path)
    return jpg_path


def _perfil_en_uso(perfil):
    for nombre in ("SingletonLock", "lockfile", "DevToolsActivePort"):
        if os.path.exists(os.path.join(perfil, nombre)):
            return True
    return False


def _lanzar(pw, perfil):
    comunes = dict(
        headless=False,
        locale="es-MX",
        no_viewport=True,
        args=["--disable-dev-shm-usage"],
    )
    extras = dict(chromium_sandbox=True, ignore_default_args=["--no-sandbox"])
    ultimo = None
    for extra in (extras, {}):
        kwargs = {**comunes, **extra}
        try:
            return pw.chromium.launch_persistent_context(
                perfil, channel="chrome", **kwargs
            )
        except TypeError:
            ultimo = None
            continue
        except Exception as e:
            ultimo = e
            try:
                return pw.chromium.launch_persistent_context(perfil, **kwargs)
            except TypeError:
                continue
            except Exception as e2:
                ultimo = e2
    if ultimo:
        raise ultimo
    return pw.chromium.launch_persistent_context(perfil, **comunes)


def _abrir_whatsapp(page):
    url = page.url or ""
    if "web.whatsapp.com" in url and url != "about:blank":
        print("WhatsApp ya estaba en esta ventana.")
        return True
    ultimo = None
    for n in range(1, 4):
        try:
            page.goto(URL_WA, wait_until="domcontentloaded", timeout=90000)
            return True
        except Exception as e:
            ultimo = e
            print(f"Intento {n}/3: no cargó WhatsApp Web.")
            page.wait_for_timeout(2000)
    print("No cargó https://web.whatsapp.com/")
    print("  1) Cierra el Chrome del experimento (la ventana de la corrida anterior) y Enter en esa consola.")
    print("  2) En el corporativo WhatsApp a veces se corta; reintenta en un rato.")
    if ultimo is not None:
        print(" ", str(ultimo).split("\n")[0])
    return False


def _hay_sesion(page):
    if page.locator('[data-testid="chat-list-search-container"]').count():
        return True
    if page.get_by_role("textbox", name="Search or start a new chat").count():
        return True
    if page.get_by_role("textbox", name="Buscar un chat o iniciar uno nuevo").count():
        return True
    if page.locator("#pane-side").count():
        return True
    return page.locator('[aria-label="Lista de chats"]').count() > 0


def _hay_qr(page):
    if page.locator("canvas[aria-label*='QR']").count():
        return True
    if page.locator("div[data-ref]").count():
        return True
    if page.get_by_text("Escanea el código QR", exact=False).count():
        return True
    return page.get_by_text("Scan the QR code", exact=False).count() > 0


def esperar_sesion(page, segundos=180):
    fin = time.time() + segundos
    aviso_qr = False
    while time.time() < fin:
        if _hay_sesion(page):
            print("Sesión lista (sin QR).")
            return True
        if _hay_qr(page) and not aviso_qr:
            print("Escanea el QR con el teléfono. Esta carpeta de Chrome se reusa la próxima vez.")
            aviso_qr = True
        page.wait_for_timeout(1000)
    print("No entró a WhatsApp a tiempo.")
    return False


def _caja_busqueda(page):
    loc = page.locator('[data-testid="chat-list-search-container"] input[data-tab="3"]')
    if loc.count():
        return loc.first
    loc = page.get_by_role("textbox", name="Search or start a new chat")
    if loc.count():
        return loc.first
    loc = page.get_by_role("textbox", name="Buscar un chat o iniciar uno nuevo")
    if loc.count():
        return loc.first
    loc = page.locator('input[aria-label="Search or start a new chat"]')
    if loc.count():
        return loc.first
    loc = page.locator('input[data-tab="3"]')
    if loc.count():
        return loc.first
    return None


def _caja_mensaje(page):
    nombres = (
        "Type a message",
        "Escribe un mensaje",
        "Mensaje",
    )
    for nombre in nombres:
        loc = page.locator(f'[contenteditable="true"][aria-label="{nombre}"]')
        if loc.count():
            return loc.last
        loc = page.get_by_role("textbox", name=nombre)
        if loc.count():
            return loc.last
    loc = page.locator("footer [contenteditable='true']")
    if loc.count():
        return loc.last
    loc = page.locator('#main footer [contenteditable="true"]')
    if loc.count():
        return loc.last
    return None


def abrir_grupo(page, nombre):
    caja = _caja_busqueda(page)
    if caja is None:
        print("No encontré la búsqueda de chats.")
        return False
    caja.click()
    page.wait_for_timeout(300)
    caja.fill(nombre)
    print(f'Busqué: {nombre}')
    page.wait_for_timeout(2000)

    chat = page.locator(f'span[title="{nombre}"]')
    if chat.count() == 0:
        chat = page.locator(f'div[title="{nombre}"]')
    if chat.count() == 0:
        chat = page.get_by_title(nombre, exact=True)
    try:
        chat.first.click(timeout=15000)
    except Exception:
        print(f'No encontré el grupo "{nombre}" en la lista filtrada.')
        return False

    try:
        page.get_by_role("banner").get_by_text(nombre, exact=False).first.wait_for(timeout=8000)
    except Exception:
        try:
            page.locator("#main header").get_by_text(nombre, exact=False).first.wait_for(timeout=8000)
        except Exception:
            page.wait_for_timeout(1500)
    print(f'Grupo abierto: {nombre}')
    return True


def _boton_enviar(page):
    for sel in (
        'div[role="button"][aria-label="Send"]',
        'div[role="button"][aria-label="Enviar"]',
        'button[aria-label="Send"]',
        'button[aria-label="Enviar"]',
        'span[data-icon="send"]',
    ):
        loc = page.locator(sel)
        if loc.count():
            return loc.last
    return page.locator("button[aria-label='Send']")


def _js_clic_enviar_media(page):
    return page.evaluate(
        """() => {
            const el = document.querySelector('div[role="button"][aria-label="Send 1 selected"]')
                || document.querySelector('div[role="button"][aria-label="Send"]')
                || document.querySelector('div[role="button"][aria-label="Enviar"]')
                || document.querySelector('[data-icon="wds-ic-send-filled"]')?.closest('[role="button"]');
            if (!el) return false;
            el.click();
            return true;
        }"""
    )


def _hay_preview_media(page):
    for sel in (
        'div[role="button"][aria-label="Send 1 selected"]',
        'div[role="button"][aria-label="Send"]',
        'div[role="button"][aria-label="Enviar"]',
    ):
        if page.locator(sel).count():
            return True
    return False


def _mantener_seleccion(page):
    txt = page.evaluate(
        """() => (document.body && document.body.innerText || '').toLowerCase()"""
    )
    if not (
        "discard selection" in txt
        or "descartar selección" in txt
        or "descartar seleccion" in txt
    ):
        return False
    print("  diálogo Discard → Cancel")
    for sel in (
        'button:has-text("Cancel")',
        'button:has-text("Cancelar")',
        'div[role="button"]:has-text("Cancel")',
        'div[role="button"]:has-text("Cancelar")',
    ):
        loc = page.locator(sel)
        if loc.count():
            try:
                loc.first.click(timeout=3000)
                page.wait_for_timeout(400)
                return True
            except Exception:
                pass
    return False


def _abrir_clip(page):
    btn = page.locator('button[aria-label="Attach"], button[aria-label="Adjuntar"]')
    if not btn.count():
        return False
    try:
        if btn.first.get_attribute("aria-expanded") == "true":
            return True
        btn.first.click(timeout=8000)
    except Exception:
        return False
    for _ in range(20):
        if btn.first.get_attribute("aria-expanded") == "true":
            return True
        if page.locator('button[role="menuitem"]').count():
            return True
        page.wait_for_timeout(200)
    return True


def _pulsar_enviar_media(page):
    _mantener_seleccion(page)
    media = page.locator('div[role="button"][aria-label="Send 1 selected"]')
    if media.count() == 0:
        media = page.locator(
            'div[role="button"][aria-label="Send"], '
            'div[role="button"][aria-label="Enviar"]'
        )
    try:
        if media.count():
            media.last.click(timeout=10000, force=True)
        elif not _js_clic_enviar_media(page):
            return False
    except Exception:
        if not _js_clic_enviar_media(page):
            return False
    page.wait_for_timeout(2000)
    _mantener_seleccion(page)
    if _hay_preview_media(page):
        if not _js_clic_enviar_media(page):
            return False
        page.wait_for_timeout(1500)
    return True


def _adjuntar_pegando(page, ruta_abs):
    """
    Pegar el JPEG en la caja del chat. WhatsApp lo trata como foto
    (no usa el input de sticker / image/* del clip).
    """
    caja = _caja_mensaje(page)
    if caja is None:
        print("  paste: no hay caja de mensaje")
        return False
    try:
        caja.click(timeout=8000)
    except Exception:
        return False
    page.wait_for_timeout(200)
    with open(ruta_abs, "rb") as f:
        b64 = __import__("base64").b64encode(f.read()).decode("ascii")
    nombre = os.path.basename(ruta_abs)
    ok = page.evaluate(
        """async ({ b64, nombre }) => {
            const caja = document.querySelector('#main footer [contenteditable="true"]')
                || document.querySelector('footer [contenteditable="true"]');
            if (!caja) return false;
            const bin = atob(b64);
            const bytes = new Uint8Array(bin.length);
            for (let i = 0; i < bin.length; i++) bytes[i] = bin.charCodeAt(i);
            const file = new File([bytes], nombre, { type: 'image/jpeg' });
            const dt = new DataTransfer();
            dt.items.add(file);
            caja.focus();
            const ev = new ClipboardEvent('paste', { bubbles: true, cancelable: true, clipboardData: dt });
            // Chromium a veces ignora clipboardData del constructor; forzar DataTransfer
            try {
                Object.defineProperty(ev, 'clipboardData', { get: () => dt });
            } catch (e) {}
            caja.dispatchEvent(ev);
            return true;
        }""",
        {"b64": b64, "nombre": nombre},
    )
    if not ok:
        print("  paste: evaluate falló")
        return False
    print("  paste JPEG en caja de mensaje")
    for _ in range(32):
        page.wait_for_timeout(250)
        _mantener_seleccion(page)
        if _hay_preview_media(page):
            print("  Preview OK tras paste.")
            return True
    print("  paste: no abrió preview")
    return False


def _descartar_preview(page):
    """Cierra un preview a medias antes de otro intento."""
    if not _hay_preview_media(page) and not page.evaluate(
        """() => {
            const t = (document.body && document.body.innerText || '').toLowerCase();
            return t.includes('discard selection') || t.includes('descartar');
        }"""
    ):
        return
    print("  cierro preview anterior...")
    page.keyboard.press("Escape")
    page.wait_for_timeout(500)
    for sel in (
        'button:has-text("Discard")',
        'button:has-text("Descartar")',
        'div[role="button"]:has-text("Discard")',
        'div[role="button"]:has-text("Descartar")',
    ):
        loc = page.locator(sel)
        if loc.count():
            try:
                loc.last.click(timeout=3000)
                page.wait_for_timeout(500)
                break
            except Exception:
                pass
    page.wait_for_timeout(400)


def _adjuntar_documento(page, ruta_abs):
    """Document: WhatsApp NUNCA lo convierte en sticker."""
    if not _abrir_clip(page):
        print("  document: no abrí Attach")
        return False
    page.wait_for_timeout(400)
    doc = page.locator(
        'button[role="menuitem"][aria-label*="Document"], '
        'button[role="menuitem"][aria-label*="Documento"]'
    )
    if not doc.count():
        print("  document: no vi menuitem Document")
        return False
    with open(ruta_abs, "rb") as f:
        raw = f.read()
    payload = {
        "name": os.path.basename(ruta_abs),
        "mimeType": "image/jpeg",
        "buffer": raw,
    }
    try:
        with page.expect_file_chooser(timeout=6000) as fc_info:
            doc.first.click(timeout=8000)
        fc_info.value.set_files(payload)
        print("  Document + file_chooser OK")
    except Exception as e:
        print("  document file_chooser:", str(e).split("\n")[0])
        infos = page.evaluate(
            """() => [...document.querySelectorAll('input[type="file"]')].map((el, i) => ({
                i, accept: el.accept || '', multiple: !!el.multiple
            }))"""
        )
        print("  inputs:", infos)
        puesto = False
        for info in infos:
            a = (info.get("accept") or "").lower().strip()
            # Solo input de documentos (*), nunca image/* ni png/webp
            if "webp" in a and "png" in a:
                continue
            if a in ("image/*",) or a.startswith("image/*"):
                continue
            if a in ("", "*", "*/*") or a.startswith("*") or (
                info.get("multiple") and "image" not in a
            ):
                try:
                    page.locator('input[type="file"]').nth(info["i"]).set_input_files(payload)
                    print(f"  set_input_files DOC input[{info['i']}] accept={a[:40]!r}")
                    puesto = True
                    break
                except Exception as e2:
                    print(f"  input[{info['i']}]:", str(e2).split("\n")[0])
        if not puesto:
            return False
    for _ in range(32):
        page.wait_for_timeout(250)
        _mantener_seleccion(page)
        if _hay_preview_media(page):
            print("  Preview OK (documento).")
            return True
    return False


def enviar_imagen(page, ruta, pie=""):
    """
    Por qué salía sticker: el input footer accept=image/* en WA actual a veces
    entra al pipeline de sticker aunque el preview se vea como foto.

    Orden:
      1) Pegar JPEG en la caja → foto real
      2) Si no, Attach → Document → nunca sticker (archivo con preview)
    """
    if not os.path.isfile(ruta):
        print("No está el JPEG de la captura.")
        return False
    ruta_abs = os.path.abspath(ruta)

    print("Intento 1: pegar imagen en el chat (foto)...")
    if _adjuntar_pegando(page, ruta_abs):
        if _pulsar_enviar_media(page):
            print("Imagen enviada (foto por paste).")
            return True
        print("  paste abrió preview pero no envió.")
    _descartar_preview(page)

    print("Intento 2: Document (garantiza que NO sea sticker)...")
    if _adjuntar_documento(page, ruta_abs):
        if _pulsar_enviar_media(page):
            print("Imagen enviada como documento (no sticker).")
            return True
        print("  documento en preview pero no envió.")

    print("No pude enviar la imagen.")
    return False


def escribir_y_enviar(page, texto):
    caja = _caja_mensaje(page)
    if caja is None:
        print("No encontré la caja de mensaje.")
        return False
    caja.click()
    page.wait_for_timeout(300)
    page.keyboard.insert_text(texto)
    page.wait_for_timeout(300)
    btn = _boton_enviar(page)
    if btn.count():
        try:
            btn.first.click(timeout=8000)
        except Exception:
            page.keyboard.press("Enter")
    else:
        page.keyboard.press("Enter")
    page.wait_for_timeout(2000)
    print("Mensaje enviado.")
    return True


def enviar_grupo(page, nombre, texto, imagen=None):
    if not abrir_grupo(page, nombre):
        return False
    if imagen:
        return enviar_imagen(page, imagen)
    return escribir_y_enviar(page, texto)


def enviar(page, telefono, texto):
    tel = "".join(c for c in telefono if c.isdigit())
    if not tel:
        print("Número vacío.")
        return False
    from urllib.parse import quote
    page.goto(
        f"https://web.whatsapp.com/send?phone={tel}&text={quote(texto)}",
        wait_until="domcontentloaded",
        timeout=120000,
    )
    if not esperar_sesion(page, 90):
        return False
    enviar_btn = page.locator('button[aria-label="Enviar"]')
    if enviar_btn.count() == 0:
        enviar_btn = page.locator('span[data-icon="send"]')
    try:
        enviar_btn.first.click(timeout=20000)
    except Exception:
        print("No encontré el botón Enviar. Revisa el número o si WhatsApp cambió la web.")
        return False
    page.wait_for_timeout(2000)
    print("Mensaje enviado (o al menos se pulsó Enviar).")
    return True


def main():
    if sys.platform == "win32":
        try:
            sys.stdout.reconfigure(encoding="utf-8")
        except Exception:
            pass

    ap = argparse.ArgumentParser(description="Consulta SQL → WhatsApp (perfil persistente).")
    ap.add_argument("--grupo", default=GRUPO_DEFAULT, help="Nombre exacto del grupo")
    ap.add_argument("--para", default="", help="Número destino con lada país, sin + (opcional)")
    ap.add_argument("--texto", default=None, help="Texto fijo; si omites, manda el resultado SQL")
    ap.add_argument("--prueba", action="store_true", help="Manda el texto de persistencia, sin SQL")
    ap.add_argument("--solo-abrir", action="store_true", help="No enviar, solo abrir sesión")
    args = ap.parse_args()

    _cargar_env()
    texto = args.texto
    imagen = None
    if args.prueba:
        texto = TEXTO_DEFAULT
    elif texto is None and not args.solo_abrir:
        try:
            texto, cols, filas = consulta_asistencia()
            xlsx = armar_excel(cols, filas)
            print("Excel:")
            print(" ", xlsx)
            imagen = captura_excel(xlsx)
        except Exception as e:
            print("No salió la consulta o la captura.")
            print(" ", str(e))
            return 1
        print("Resultado SQL:")
        print(texto)
        print()

    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        print("Falta Playwright. En la carpeta del repo:")
        print("  python -m pip install -r requirements.txt")
        print("  python -m playwright install chromium")
        return 1

    perfil = os.path.join(_dir_script(), PERFIL)
    os.makedirs(perfil, exist_ok=True)
    print("Perfil Chrome (no es tu Chrome de diario):")
    print(" ", perfil)
    if _perfil_en_uso(perfil):
        print("Ese perfil parece ocupado. Cierra el Chrome del experimento (corrida anterior) y vuelve a correr.")

    with sync_playwright() as pw:
        try:
            ctx = _lanzar(pw, perfil)
        except Exception as e:
            print("No se abrió Chrome.")
            print("  Cierra la ventana Chrome del experimento si sigue abierta.")
            print("  python -m pip install -r requirements.txt")
            print("  python -m playwright install chromium")
            print(" ", str(e).split("\n")[0])
            return 1
        page = ctx.pages[0] if ctx.pages else ctx.new_page()
        if not _abrir_whatsapp(page):
            ctx.close()
            return 2
        ok = esperar_sesion(page)
        if ok and not args.solo_abrir:
            try:
                if args.para:
                    ok = enviar(page, args.para, texto)
                else:
                    ok = enviar_grupo(page, args.grupo, texto, imagen=imagen)
            except Exception as e:
                print("Falló el envío (Chrome sigue abierto).")
                print(" ", str(e).split("\n")[0])
                ok = False
        print("Cierra la ventana de Chrome cuando termines (o Enter aquí).")
        try:
            input()
        except EOFError:
            page.wait_for_timeout(8000)
        ctx.close()
    return 0 if ok else 2


if __name__ == "__main__":
    raise SystemExit(main())
